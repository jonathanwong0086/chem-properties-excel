import json
import tempfile
import unittest
from pathlib import Path

import openpyxl

from generate_excel import FIELDS, build_excel, lookup_regulated_chemical


def _write_index(path: Path) -> None:
    data = {
        "schema_version": "1.0",
        "sources": [{"list_type": "易制毒"}, {"list_type": "易制爆"}],
        "records": [
            {
                "record_id": "acetone",
                "list_type": "易制毒",
                "name": "丙酮",
                "aliases": ["二甲基酮"],
                "cas_no": ["67-64-1"],
                "category": "第三类",
                "item_no": "3-1",
                "scope_condition": "",
                "salt_rule": False,
                "requires_condition_check": False,
                "source_locator": "条例附表",
            },
            {
                "record_id": "peroxide",
                "list_type": "易制爆",
                "name": "过氧化氢",
                "aliases": ["双氧水"],
                "cas_no": ["7722-84-1"],
                "category": "氧化剂",
                "item_no": "5.2",
                "scope_condition": "含量不低于20%",
                "hazard_classification": "氧化性液体",
                "requires_condition_check": True,
                "source_row": "原目录第5.2项",
            },
        ],
        "search_index": {
            "by_cas": {"67-64-1": ["acetone"], "7722-84-1": ["peroxide"]},
            "by_name": {"丙酮": ["acetone"], "过氧化氢": ["peroxide"]},
            "by_alias": {
                "二甲基酮": ["acetone"],
                "双氧水": ["peroxide"],
            },
        },
    }
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")


def _chemicals():
    return [
        {
            "name": "丙酮", "cas_no": "67-64-1", "fire": "甲类", "mp": "-94.7", "toxic_gas": "否",
            "gbt_toxicity": "未列入GB/T 42594-2023",
            "gbt_combustion": "未列入GB/T 42594-2023",
            "gbt_compatibility": "未列入GB/T 42594-2023",
            "gbt_reactivity": "未列入GB/T 42594-2023",
            "gbt_hazard": "未列入GB/T 42594-2023",
            "gbt_ghs": "未列入GB/T 42594-2023",
            "sources": {
                "cas_no": "PubChem https://example.test/acetone",
                "mp": "NIST https://example.test/mp",
            },
        },
        {
            "name": "过氧化氢", "cas_no": "7722-84-1", "fire": "非可燃",
            "gbt_toxicity": "III", "gbt_combustion": "氧化性介质",
            "gbt_compatibility": "避免铜", "gbt_reactivity": "受热分解",
            "gbt_hazard": "强氧化性", "gbt_ghs": "氧化性液体",
        },
    ]


class GenerateExcelTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.directory = Path(self.temp.name)
        self.index_path = self.directory / "index.json"
        _write_index(self.index_path)

    def tearDown(self):
        self.temp.cleanup()

    def test_lookup_uses_cas_alias_and_condition(self):
        index = json.loads(self.index_path.read_text(encoding="utf-8"))
        acetone = lookup_regulated_chemical(index, "错误名称", "67-64-1")
        peroxide = lookup_regulated_chemical(index, "双氧水", "")
        self.assertEqual(acetone["is_precursor"], "是（第三类）")
        self.assertIn("易制毒化学品管理条例", acetone["sources"]["is_precursor"])
        self.assertEqual(peroxide["is_explosive_precursor"], "需核实浓度/形态")
        self.assertIn("含量不低于20%", peroxide["sources"]["is_explosive_precursor"])
        self.assertNotIn("匹配依据", peroxide["sources"]["is_precursor"])

    def test_missing_index_never_reports_a_false_negative(self):
        result = lookup_regulated_chemical(None, "丙酮", "67-64-1")
        self.assertEqual(result["is_precursor"], "未核实（缺少目录索引）")
        self.assertEqual(result["is_explosive_precursor"], "未核实（缺少目录索引）")
        with self.assertRaises(FileNotFoundError):
            build_excel(
                _chemicals(), "缺索引测试", self.directory,
                regulatory_index_path=self.directory / "missing.json",
            )

    def test_workbook_contract_comments_merge_and_safe_name(self):
        project = "很长的项目名称用于验证工作表名称截断且不会超出Excel限制/含非法字符"
        first = build_excel(_chemicals(), project, self.directory, regulatory_index_path=self.index_path)
        second = build_excel(_chemicals(), project, self.directory, regulatory_index_path=self.index_path)
        self.assertNotEqual(first, second)
        self.assertTrue(Path(first).exists() and Path(second).exists())
        self.assertTrue(Path(second).stem.endswith("_v2"))
        self.assertNotIn("/", Path(first).name)
        workbook = openpyxl.load_workbook(first)
        try:
            sheet = workbook.worksheets[0]
            self.assertLessEqual(len(sheet.title), 31)
            self.assertEqual(sheet.freeze_panes, "C3")
            self.assertEqual(
                [sheet.cell(2, column).value for column in range(1, 27)],
                [heading for heading, _ in FIELDS],
            )
            self.assertEqual(sheet.max_column, 26)
            self.assertEqual(sheet["C3"].comment.text, "数据来源:\nPubChem https://example.test/acetone")
            self.assertEqual(sheet["Y3"].value, "是（第三类）")
            self.assertIn("易制毒化学品管理条例", sheet["Y3"].comment.text)
            self.assertEqual(sheet["Z4"].value, "需核实浓度/形态")
            self.assertIn("含量不低于20%", sheet["Z4"].comment.text)
            merged = {str(item) for item in sheet.merged_cells.ranges}
            self.assertIn("S3:X3", merged)
            self.assertEqual(sheet["S3"].value, "本地参考未命中，需核对GB/T 42594-2023原文")
            self.assertNotIn("S4:X4", merged)
            self.assertEqual(workbook.worksheets[1].title, "数据来源汇总")
            self.assertEqual(workbook.worksheets[1].max_row, 20)
        finally:
            workbook.close()

    def test_optional_sheets_append_regulatory_columns(self):
        chemical_list = [{"name": "丙酮", "cas_no": "67-64-1", "category": "溶剂"}]
        equipment = [
            {"name": "丙酮", "cas_no": "67-64-1", "category": "溶剂", "equipment_tag": "V-1"},
            {"name": "丙酮", "cas_no": "67-64-1", "category": "溶剂", "equipment_tag": "V-2"},
        ]
        output = build_excel(
            _chemicals(), "测试", self.directory,
            chemical_list=chemical_list, equipment_mapping=equipment,
            regulatory_index_path=self.index_path,
        )
        workbook = openpyxl.load_workbook(output)
        try:
            self.assertEqual(
                workbook.sheetnames,
                ["测试化学品物性数据", "数据来源汇总", "化学品清单", "介质-设备对照表"],
            )
            listing, equipment_sheet = workbook["化学品清单"], workbook["介质-设备对照表"]
            self.assertEqual(listing.cell(2, 9).value, "是否易制毒")
            self.assertEqual(listing.cell(2, 10).value, "是否易制爆")
            self.assertEqual(listing.cell(3, 9).value, "是（第三类）")
            self.assertEqual(equipment_sheet.cell(2, 10).value, "是否易制毒")
            self.assertEqual(equipment_sheet.cell(2, 11).value, "是否易制爆")
            self.assertEqual(equipment_sheet.cell(3, 10).value, "是（第三类）")
            merged = {str(item) for item in equipment_sheet.merged_cells.ranges}
            self.assertTrue({"A3:A4", "B3:B4", "C3:C4"}.issubset(merged))
        finally:
            workbook.close()

    def test_missing_fire_class_does_not_break_generation(self):
        output = build_excel(
            [{"name": "高锰酸钾", "cas_no": "7722-64-7"}],
            "缺少火灾类别",
            self.directory,
            regulatory_index_path=self.index_path,
        )
        workbook = openpyxl.load_workbook(output)
        try:
            self.assertIsNone(workbook.worksheets[0]["I3"].value)
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
